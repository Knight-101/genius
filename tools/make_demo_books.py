"""Writes the demo ledgers: three invented businesses, three years of double-entry books each (April 2023 - March 2026), in the file
layouts Indian small businesses actually export. Everything is invented; no real company data is used.

  Kaveri Foods Pvt Ltd        food processing, small (plant above the micro limit), steady, collects on time      -> Zoho-style journal (.xlsx)
  Sharma Traders Pvt Ltd      building-materials distributor, small, two big customers paying late since mid-2025 -> Tally day book (.csv)
  Swift Freight Carriers      logistics, micro, lost volume in 2025, cash credit maxed, statutory dues slipping    -> messy journal (.xlsx)
  Sharma Traders (by account) the same books as a ledger-by-account export, to show a second layout                -> .xlsx
"""
from __future__ import annotations

import datetime as dt
import math
import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "samples"
OUT.mkdir(exist_ok=True)
START, CUT, END = dt.date(2023, 1, 1), dt.date(2023, 4, 1), dt.date(2026, 3, 31)


def L(acct, dr=0.0, cr=0.0, party=None, ref=None, due=None, narr=None):
    return (acct, float(dr), float(cr), party, ref, due, narr)


class Book:
    def __init__(self):
        self.lines, self.n = [], 0

    def post(self, d, vtype, lines):
        self.n += 1
        td = round(sum(l[1] for l in lines), 2); tc = round(sum(l[2] for l in lines), 2)
        assert abs(td - tc) < 0.011, (vtype, d, td, tc)
        for acct, dr, cr, party, ref, due, narr in lines:
            self.lines.append(dict(date=pd.Timestamp(d), eid=self.n, vtype=vtype, account=acct, debit=round(dr, 2), credit=round(cr, 2),
                                   party=party, invoice_ref=ref, due_date=pd.Timestamp(due) if due is not None else None, narration=narr))


def month_end(y, m):
    return (pd.Timestamp(year=y, month=m, day=1) + pd.offsets.MonthEnd(0)).date()


def dom(y, m, day):
    return dt.date(y, m, min(day, month_end(y, m).day))


def weekday(rng, y, m):
    while True:
        d = dom(y, m, int(rng.integers(1, 29)))
        if d.weekday() < 6:
            return d


def split(rng, total, n):
    if n == 1:
        return [round(total, 2)]
    w = rng.dirichlet(np.ones(n) * 4)
    return [round(total * x, 2) for x in w]


def simulate(spec: dict, seed: int):
    rng = np.random.default_rng(seed)
    bk = Book()
    B, CC, TL = spec["bank"], spec["cc"], spec.get("tl")
    REC, PAY, INV = spec["recv_acct"], spec["pay_acct"], spec.get("inv_acct")
    opening = [L(B, dr=spec["cash0"])]
    for a, v in spec["fa0"].items():
        opening.append(L(a, dr=v))
    if INV:
        opening.append(L(INV, dr=spec["inv0"]))
    eq = sum(l[1] for l in opening)
    if TL and spec.get("tl0"):
        opening.append(L(TL, cr=spec["tl0"])); eq -= spec["tl0"]
    opening.append(L(spec["capital"], cr=round(eq, 2)))
    bk.post(START, "Opening", opening)
    months = pd.period_range(pd.Timestamp(START), pd.Timestamp(END), freq="M")
    out_gst, in_gst, seq, seqb = {}, {}, {}, 0
    tl_bal = spec.get("tl0", 0.0)
    for t, per in enumerate(months):
        y, m = per.year, per.month
        fy = y if m >= 4 else y - 1
        mult = spec["season"][m - 1] * (1 + spec["growth"]) ** (t / 12) * rng.normal(1, spec.get("noise", .03))
        for (a, b), k in spec.get("shocks", {}).items():
            if pd.Period(a, "M") <= per <= pd.Period(b, "M"):
                mult *= k
        target = spec["sales0"] * mult
        custs = spec["customers"]
        w = np.array([c[1] for c in custs]) * rng.normal(1, .08, len(custs))
        w = np.clip(w, 0.001, None); w /= w.sum()
        cogs_m = sales_m = 0.0
        for (name, _, terms, lm, ls), share in zip(custs, w):
            n_inv = max(1, int(round(target * share / spec["ticket"])))
            for part in split(rng, target * share, min(n_inv, 12)):
                d = weekday(rng, y, m)
                net = round(part, 2); gst = round(net * spec["gst"], 2); gross = round(net + gst, 2)
                seq[fy] = seq.get(fy, 0) + 1
                ref = f"{spec['prefix']}/{str(fy)[2:]}-{str(fy + 1)[2:]}/{seq[fy]:04d}"
                due = d + dt.timedelta(days=terms)
                bk.post(d, "Sales Invoice", [L(REC, dr=gross, party=name, ref=ref, due=due, narr=f"Invoice {ref}"),
                                             L(spec["sales_acct"], cr=net), L(spec["gst_out"], cr=gst)])
                if INV:
                    cost = round(net * (1 - spec["gm"]) * rng.normal(1, .01), 2)
                    bk.post(d, "Cost of Sales", [L(spec["cogs_acct"], dr=cost), L(INV, cr=cost)])
                    cogs_m += cost
                sales_m += net
                out_gst[(y, m)] = round(out_gst.get((y, m), 0) + gst, 2)
                mu, sd = lm, ls
                for (pn, a, b), extra in spec.get("late", {}).items():
                    if pn == name and pd.Timestamp(a) <= pd.Timestamp(d) <= pd.Timestamp(b):
                        mu += extra
                rd = due + dt.timedelta(days=int(max(round(rng.normal(mu, sd)), -3)))
                rd = max(rd, d + dt.timedelta(days=1))
                if rd <= END:
                    bk.post(rd, "Receipt", [L(B, dr=gross), L(REC, cr=gross, party=name, ref=ref, narr=f"Receipt against {ref}")])
        # purchases
        buy = cogs_m * rng.normal(1.0, .03) if INV else sales_m * (1 - spec["gm"])
        for sname, sw, sterms, slag, _msme in spec["suppliers"]:
            for part in split(rng, buy * sw, 2):
                d = weekday(rng, y, m)
                net = round(part, 2); gst = round(net * spec["gst"], 2); gross = round(net + gst, 2)
                seqb += 1
                ref = f"BILL/{seqb:05d}"
                due = d + dt.timedelta(days=sterms)
                debit_acct = INV if INV else spec["cogs_acct"]
                bk.post(d, "Purchase Bill", [L(debit_acct, dr=net), L(spec["gst_in"], dr=gst), L(PAY, cr=gross, party=sname, ref=ref, due=due, narr=f"Bill {ref}")])
                in_gst[(y, m)] = round(in_gst.get((y, m), 0) + gst, 2)
                lag = slag + int(round(rng.normal(0, 3)))
                for (a, b), extra in spec.get("stretch", {}).items():
                    if pd.Timestamp(a) <= pd.Timestamp(due) <= pd.Timestamp(b):
                        lag += extra
                pdt = max(due + dt.timedelta(days=lag), d + dt.timedelta(days=1))
                if pdt <= END:
                    bk.post(pdt, "Payment", [L(PAY, dr=gross, party=sname, ref=ref, narr=f"Payment {ref}"), L(B, cr=gross)])
        me = month_end(y, m)
        sal = spec["salary"] * (1.07 ** (t // 12))
        pf, tds = round(sal * .06, 2), round(sal * .03, 2)
        bk.post(me, "Payroll", [L("Salaries and Wages", dr=sal), L(B, cr=round(sal - pf - tds, 2)), L("PF and ESI Payable", cr=pf), L("TDS Payable", cr=tds)])
        ny, nm = (y + (m == 12), m % 12 + 1)
        delay = spec.get("stat_delay", {}).get(str(per), 0)
        sd_ = dom(ny, nm, 15) + dt.timedelta(days=delay)
        if sd_ <= END:
            bk.post(sd_, "Payment", [L("PF and ESI Payable", dr=pf), L("TDS Payable", dr=tds), L(B, cr=round(pf + tds, 2), narr="Statutory remittance")])
        for acct, amt, day in spec["fixed"]:
            v = round(amt * (1.05 ** (t // 12)) * (rng.uniform(.9, 1.1) if acct == "Electricity Charges" else 1), 2)
            bk.post(dom(y, m, day), "Payment", [L(acct, dr=v), L(B, cr=v)])
        fr = round(sales_m * spec.get("freight", 0), 2)
        if fr > 0:
            bk.post(dom(y, m, 27), "Payment", [L("Freight Outward", dr=fr), L(B, cr=fr)])
        if m == 10:
            bk.post(dom(y, m, 14), "Payment", [L("Insurance", dr=spec["insurance"]), L(B, cr=spec["insurance"])])
        drw = spec["drawings"]
        bk.post(dom(y, m, 5), "Payment", [L("Owner's Drawings", dr=drw, narr="Owner drawings"), L(B, cr=drw)])
        bk.post(me, "Journal", [L("Depreciation", dr=spec["dep"]), L("Accumulated Depreciation", cr=spec["dep"])])
        py_, pm_ = (y - (m == 1), 12 if m == 1 else m - 1)
        if (py_, pm_) in out_gst and dom(y, m, 20) <= END:
            o, i_ = out_gst[(py_, pm_)], in_gst.get((py_, pm_), 0)
            if o >= i_:
                bk.post(dom(y, m, 20), "Payment", [L(spec["gst_out"], dr=o), L(spec["gst_in"], cr=i_), L(B, cr=round(o - i_, 2), narr="GST remittance")])
        if TL and tl_bal > 1:
            it = round(tl_bal * spec["tl_rate"] / 12, 2)
            pr = round(min(tl_bal, spec["emi"] - it), 2)
            tl_bal -= pr
            bk.post(dom(y, m, 7), "Payment", [L(TL, dr=pr), L("Interest on Term Loan", dr=it), L(B, cr=round(pr + it, 2), narr="EMI")])
        for when, amt in spec.get("capital_in", {}).items():
            if str(per) == when:
                bk.post(dom(y, m, 18), "Receipt", [L(B, dr=amt), L(spec["capital"], cr=amt, narr="Capital introduced by owner")])
    # daily cash-credit sweep and interest
    df = pd.DataFrame(bk.lines)
    delta = df[df.account == B].groupby("date").apply(lambda g: g.debit.sum() - g.credit.sum(), include_groups=False).to_dict()
    bank = cc = 0.0
    MIN, TARGET, HIGH, LIMIT = spec["cc_policy"]
    for d in pd.date_range(START, END):
        bank += delta.get(d, 0.0)
        if bank < MIN and cc < LIMIT:
            draw = min(LIMIT - cc, math.ceil((TARGET - bank) / 100_000) * 100_000)
            if draw > 0:
                bk.post(d.date(), "Bank Sweep", [L(B, dr=draw), L(CC, cr=draw, narr="Cash credit drawdown")])
                bank += draw; cc += draw
        elif bank > HIGH and cc > 0:
            rep = min(cc, math.floor((bank - (TARGET + HIGH) / 2) / 100_000) * 100_000)
            if rep > 0:
                bk.post(d.date(), "Bank Sweep", [L(CC, dr=rep, narr="Cash credit repayment"), L(B, cr=rep)])
                bank -= rep; cc -= rep
        if d == pd.Timestamp(month_end(d.year, d.month)) and cc > 0:
            it = round(cc * spec["cc_rate"] / 12, 2)
            bk.post(d.date(), "Journal", [L("Interest on Cash Credit", dr=it), L(CC, cr=it, narr="Interest debited to CC")])
            cc += it
    return bk


def export(bk: Book, spec: dict) -> pd.DataFrame:
    """Books from CUT onwards, with balances at CUT carried in as opening lines (per invoice for the party ledgers)."""
    df = pd.DataFrame(bk.lines)
    pre, post = df[df.date < pd.Timestamp(CUT)], df[df.date >= pd.Timestamp(CUT)].copy()
    lines = []
    party_ledgers = (spec["recv_acct"], spec["pay_acct"])
    for acct, g in pre.groupby("account"):
        if acct in party_ledgers:
            for (party, ref, due), gg in g.groupby(["party", "invoice_ref", "due_date"], dropna=False):
                net = round(gg.debit.sum() - gg.credit.sum(), 2)
                if abs(net) > 0.004:
                    lines.append(dict(account=acct, debit=max(net, 0), credit=max(-net, 0), party=party, invoice_ref=ref, due_date=due))
        else:
            net = round(g.debit.sum() - g.credit.sum(), 2)
            if abs(net) > 0.004:
                lines.append(dict(account=acct, debit=max(net, 0), credit=max(-net, 0), party=None, invoice_ref=None, due_date=None))
    keep = [l for l in lines if l["account"] in spec["bs_accounts"]]
    plug = round(sum(l["debit"] for l in keep) - sum(l["credit"] for l in keep), 2)
    keep.append(dict(account=spec["capital"], debit=max(-plug, 0), credit=max(plug, 0), party=None, invoice_ref=None, due_date=None))
    op = pd.DataFrame(keep)
    op["date"], op["eid"], op["vtype"], op["narration"] = pd.Timestamp(CUT), 0, "Opening balance", "Opening balance as on 01-Apr-2023"
    allp = pd.concat([op, post], ignore_index=True).sort_values(["date", "eid"], kind="stable").reset_index(drop=True)
    ids = {e: f"V{i + 1:05d}" for i, e in enumerate(dict.fromkeys(allp.eid))}
    allp["entry_no"] = allp.eid.map(ids)
    assert abs(allp.debit.sum() - allp.credit.sum()) < 0.05
    return allp[["date", "entry_no", "vtype", "account", "debit", "credit", "party", "invoice_ref", "due_date", "narration"]]


def indian(x):
    s = f"{abs(x):.2f}"
    ip, fp = s.split(".")
    if len(ip) > 3:
        head, tail = ip[:-3], ip[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:]); head = head[:-2]
        if head:
            parts.insert(0, head)
        ip = ",".join(parts) + "," + tail
    return f"{ip}.{fp}"


def write_tally(df, path, company, party_accounts):
    import csv
    rows = [["Day Book"], [company], ["1-Apr-23 to 31-Mar-26"], [], ["Date", "Particulars", "Vch Type", "Vch No.", "Ref No.", "Due Date", "Debit", "Credit"]]
    for eno, g in df.groupby("entry_no", sort=False):
        first = True
        for r in g.itertuples():
            name = r.party if (r.account in party_accounts and isinstance(r.party, str)) else r.account
            rows.append([r.date.strftime("%d-%m-%Y") if first else "", name, r.vtype if first else "", eno if first else "", r.invoice_ref or "",
                         r.due_date.strftime("%d-%m-%Y") if pd.notna(r.due_date) else "", indian(r.debit) if r.debit else "", indian(r.credit) if r.credit else ""])
            first = False
    rows.append(["", "Grand Total", "", "", "", "", indian(df.debit.sum()), indian(df.credit.sum())])
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        csv.writer(fh).writerows(rows)


def write_zoho(df, path, company):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Journal Report"
    ws.append(["Journal Report"]); ws.append([company]); ws.append(["From 01/04/2023 To 31/03/2026"]); ws.append([])
    ws.append(["Journal Date", "Journal #", "Account", "Debit", "Credit", "Contact Name", "Reference#", "Due Date", "Notes"])
    for c in ws[5]:
        c.font = Font(bold=True)
    for r in df.itertuples():
        ws.append([r.date.to_pydatetime(), r.entry_no, r.account, r.debit or None, r.credit or None, r.party if isinstance(r.party, str) else None,
                   r.invoice_ref if isinstance(r.invoice_ref, str) else None, r.due_date.to_pydatetime() if pd.notna(r.due_date) else None,
                   r.narration if isinstance(r.narration, str) else None])
    ws.append([None, None, "Total", round(df.debit.sum(), 2), round(df.credit.sum(), 2)])
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row):
        row[0].number_format = "dd/mm/yyyy"; row[7].number_format = "dd/mm/yyyy"
    wb.save(path)


def write_messy(df, path, company):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Journal"
    ws.append([company + " - Journal Vouchers"]); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    ws.append(["Currency: INR (amounts in Rs.)"]); ws.append([])
    ws.append(["Txn Date", "JV No", "Ledger Name", "Amount", "Dr/Cr", "Customer/Vendor", "Bill No", "Bill Due", "Description"])
    for c in ws[4]:
        c.font = Font(bold=True)
    cur = None
    for r in df.itertuples():
        ym = (r.date.year, r.date.month)
        if cur is not None and ym != cur:
            ws.append([None, None, f"Subtotal {pd.Timestamp(year=cur[0], month=cur[1], day=1).strftime('%b-%y')}", None, None])
        cur = ym
        amt = r.debit or r.credit
        ws.append([r.date.strftime("%d %b %Y"), r.entry_no, r.account, f"Rs. {indian(amt)}", "Dr" if r.debit else "Cr",
                   r.party if isinstance(r.party, str) else None, r.invoice_ref if isinstance(r.invoice_ref, str) else None,
                   r.due_date.strftime("%d %b %Y") if pd.notna(r.due_date) else None, r.narration if isinstance(r.narration, str) else None])
    ws.append([None, None, "Total"])
    n = wb.create_sheet("Notes"); n.append(["Prepared by accounts team. Figures unaudited."])
    wb.save(path)


def write_by_account(df, path, company, party_accounts):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "General Ledger"
    ws.append([company]); ws.append(["General Ledger - 1-Apr-23 to 31-Mar-26"]); ws.append([])
    work = df.copy()
    work["ledger"] = [r.party if (r.account in party_accounts and isinstance(r.party, str)) else r.account for r in work.itertuples()]
    work["group"] = [("Sundry Debtors" if r.account == party_accounts[0] else "Sundry Creditors") if r.account in party_accounts else "" for r in work.itertuples()]
    by_entry = {k: g for k, g in work.groupby("entry_no")}
    for ledger, g in work.groupby("ledger", sort=True):
        ws.append([f"Account: {ledger}" + (f" | Group: {g.group.iloc[0]}" if g.group.iloc[0] else "")])
        ws.append(["Date", "Particulars", "Vch Type", "Vch No.", "Ref No.", "Due Date", "Debit", "Credit", "Balance"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        bal = 0.0
        opening = g[g.vtype == "Opening balance"]
        for r in g.itertuples():
            ent = by_entry[r.entry_no]
            others = ent[ent.ledger != ledger]
            counter = "Opening Balance" if r.vtype == "Opening balance" else (others.loc[(others.debit + others.credit).idxmax(), "ledger"] if len(others) else "")
            bal += r.debit - r.credit
            ws.append([r.date.to_pydatetime(), counter, r.vtype, r.entry_no, r.invoice_ref if isinstance(r.invoice_ref, str) else None,
                       r.due_date.to_pydatetime() if pd.notna(r.due_date) else None, r.debit or None, r.credit or None, f"{indian(bal)} {'Dr' if bal >= 0 else 'Cr'}"])
            ws.cell(row=ws.max_row, column=1).number_format = "dd-mmm-yy"
        ws.append([None, "Closing Balance", None, None, None, None, None, None, f"{indian(bal)} {'Dr' if bal >= 0 else 'Cr'}"])
        ws.append([])
    wb.save(path)


COMMON_BS = {"HDFC Current A/c", "ICICI Current A/c", "SBI Current A/c", "Sundry Debtors", "Sundry Creditors", "Inventory - Trading Goods",
             "Stock - Raw Materials and Finished Goods", "GST Input Credit", "GST Output Tax", "Plant and Machinery", "Furniture and Fixtures",
             "Vehicles", "Accumulated Depreciation", "HDFC CC A/c", "ICICI CC A/c", "SBI CC A/c", "HDFC Term Loan", "SBI Term Loan",
             "ICICI Term Loan", "PF and ESI Payable", "TDS Payable", "Capital Account", "Owner's Drawings", "Factory Building", "Trucks and Trailers",
             "Computers"}

SPECS = {
    "sharma": dict(
        company="Sharma Traders Pvt Ltd", sector="Trading · building materials", prefix="ST", bank="HDFC Current A/c", cc="HDFC CC A/c", tl="HDFC Term Loan",
        recv_acct="Sundry Debtors", pay_acct="Sundry Creditors", inv_acct="Inventory - Trading Goods", sales_acct="Sales - Building Materials",
        cogs_acct="Cost of Goods Sold", gst_out="GST Output Tax", gst_in="GST Input Credit", capital="Capital Account", gst=0.18,
        cash0=3_500_000, fa0={"Furniture and Fixtures": 1_800_000, "Vehicles": 6_500_000, "Plant and Machinery": 9_000_000}, inv0=52_000_000,
        tl0=16_000_000, tl_rate=0.105, emi=340_000,
        sales0=29_000_000, growth=0.07, noise=.04, gm=0.13, ticket=400_000,
        season=[1.05, 1.08, 1.15, 1.08, 1.05, .90, .82, .85, .92, 1.0, 1.0, 1.10],
        customers=[("Gupta Hardware", .22, 30, 12, 6), ("Mehta Build Solutions", .12, 30, 15, 6), ("Om Sai Constructions", .07, 30, 10, 5),
                   ("Krishna Cement Agencies", .06, 30, 8, 4), ("Patel Infra Projects", .06, 45, 12, 6), ("Rao Buildcon", .05, 30, 9, 4),
                   ("Joshi Hardware Mart", .05, 30, 6, 3), ("Verma Sanitary Stores", .05, 30, 14, 6), ("Nair Builders", .04, 30, 7, 3),
                   ("Reddy Tiles and Stones", .04, 30, 10, 4), ("Khanna Plumbing Supplies", .04, 30, 8, 3), ("Iyer Steel Fabricators", .04, 30, 9, 4),
                   ("Bansal Paints", .04, 30, 6, 3), ("Desai Home Centre", .04, 30, 11, 4), ("Saxena Electricals", .04, 30, 7, 3),
                   ("Thakur Contractors", .04, 30, 20, 8)],
        late={("Gupta Hardware", "2025-06-01", "2026-03-31"): 38, ("Mehta Build Solutions", "2025-08-01", "2026-03-31"): 30},
        suppliers=[("Rajan Polymers", .30, 45, 2, True), ("Sundaram Industries", .25, 45, 3, True), ("UltraBuild Cement Ltd", .25, 30, 1, False),
                   ("Lakshmi Plastics", .10, 30, 4, True), ("Apex Steel Distributors", .10, 45, 2, False)],
        stretch={("2025-10-01", "2026-03-31"): 9},
        salary=1_500_000, fixed=[("Rent - Warehouse", 320_000, 1), ("Electricity Charges", 70_000, 10), ("Telephone and Internet", 12_000, 12),
                                  ("Professional Fees", 45_000, 25), ("Advertising Expenses", 40_000, 20), ("Bank Charges", 15_000, 28)],
        freight=0.012, insurance=150_000, drawings=850_000, dep=180_000,
        cc_policy=(5_000_000, 8_000_000, 14_000_000, 25_000_000), cc_rate=0.108, inputs=dict(cc_limit=25_000_000)),
    "kaveri": dict(
        company="Kaveri Foods Pvt Ltd", sector="Manufacturing · food processing", prefix="KF", bank="SBI Current A/c", cc="SBI CC A/c", tl="SBI Term Loan",
        recv_acct="Sundry Debtors", pay_acct="Sundry Creditors", inv_acct="Stock - Raw Materials and Finished Goods", sales_acct="Sales - Packaged Foods",
        cogs_acct="Cost of Goods Sold", gst_out="GST Output Tax", gst_in="GST Input Credit", capital="Capital Account", gst=0.12,
        cash0=9_000_000, fa0={"Plant and Machinery": 32_000_000, "Factory Building": 15_000_000, "Furniture and Fixtures": 800_000}, inv0=11_000_000,
        tl0=6_000_000, tl_rate=0.098, emi=260_000,
        sales0=10_500_000, growth=0.11, noise=.03, gm=0.30, ticket=350_000,
        season=[1.0, .96, 1.02, .96, .95, .97, .98, 1.0, 1.04, 1.12, 1.08, .98],
        customers=[("Annapurna Retail", .16, 30, 2, 2), ("FreshMart Chain", .14, 30, 3, 2), ("Metro Wholesale Foods", .12, 30, 1, 2),
                   ("Sai Provision Stores", .08, 21, 2, 2), ("Hotel Supply Co", .08, 30, 4, 3), ("Daily Needs Distributors", .08, 30, 2, 2),
                   ("Green Basket", .07, 30, 3, 2), ("Spice Route Exports Agents", .07, 30, 1, 1), ("Royal Caterers", .07, 30, 5, 3),
                   ("Bharat Kirana Network", .07, 21, 2, 2), ("Coastal Foods Depot", .06, 30, 3, 2)],
        suppliers=[("Cauvery Agro Producers", .40, 30, 0, True), ("Deccan Packaging", .20, 30, 1, True), ("Nandi Oils", .25, 30, 0, False),
                   ("Malnad Spices", .15, 30, 0, True)],
        salary=1_250_000, fixed=[("Rent - Godown", 90_000, 1), ("Electricity Charges", 260_000, 10), ("Telephone and Internet", 10_000, 12),
                                  ("Professional Fees", 30_000, 25), ("Repairs and Maintenance", 60_000, 18), ("Bank Charges", 6_000, 28)],
        freight=0.015, insurance=300_000, drawings=250_000, dep=420_000,
        cc_policy=(2_500_000, 5_000_000, 9_000_000, 10_000_000), cc_rate=0.098, inputs=dict(cc_limit=10_000_000)),
    "swift": dict(
        company="Swift Freight Carriers", sector="Services · logistics", prefix="SF", bank="ICICI Current A/c", cc="ICICI CC A/c", tl="ICICI Term Loan",
        recv_acct="Sundry Debtors", pay_acct="Sundry Creditors", inv_acct=None, sales_acct="Freight Income",
        cogs_acct="Direct Costs - Fuel and Vehicle Hire", gst_out="GST Output Tax", gst_in="GST Input Credit", capital="Capital Account", gst=0.12,
        cash0=1_200_000, fa0={"Trucks and Trailers": 9_000_000, "Computers": 300_000}, inv0=0,
        tl0=6_500_000, tl_rate=0.12, emi=230_000,
        sales0=6_400_000, growth=0.02, noise=.05, gm=0.27, ticket=250_000,
        season=[.98, .97, 1.04, 1.0, .99, .92, .92, .97, 1.02, 1.10, 1.09, 1.0],
        shocks={("2025-04", "2025-09"): 0.68, ("2025-10", "2026-03"): 0.74},
        customers=[("Prime Retail Logistics", .34, 30, 18, 8), ("Vega Auto Parts", .14, 45, 25, 10), ("Coastal Agro Exports", .10, 30, 15, 6),
                   ("Metro Pharma Distributors", .09, 30, 10, 5), ("Sunrise Textiles", .08, 45, 30, 12), ("Kumar Steel Rolling", .08, 30, 22, 8),
                   ("Arihant Electronics", .07, 30, 12, 5), ("Shree Cement Depot", .05, 30, 18, 6), ("Bharat Paper Mills", .05, 30, 20, 8)],
        late={("Prime Retail Logistics", "2025-05-01", "2026-03-31"): 45, ("Vega Auto Parts", "2025-07-01", "2026-03-31"): 30},
        suppliers=[("Highway Fuel Station", .45, 15, 3, False), ("Rathore Fleet Hire", .30, 30, 6, True), ("Toll and Tyre Services", .25, 30, 5, True)],
        stretch={("2025-06-01", "2026-03-31"): 28},
        stat_delay={"2025-07": 25, "2025-08": 40, "2025-09": 55, "2025-10": 60, "2025-11": 70, "2025-12": 80, "2026-01": 75, "2026-02": 40},
        salary=850_000, fixed=[("Office Rent", 70_000, 1), ("Electricity Charges", 15_000, 10), ("Telephone and Internet", 9_000, 12),
                                ("Professional Fees", 20_000, 25), ("Repairs and Maintenance", 90_000, 18), ("Bank Charges", 5_000, 28)],
        freight=0.0, insurance=240_000, drawings=260_000, dep=160_000, capital_in={"2025-11": 1_000_000},
        cc_policy=(400_000, 900_000, 2_500_000, 3_000_000), cc_rate=0.125, inputs=dict(cc_limit=3_000_000)),
}


def build(which=("sharma", "kaveri", "swift"), seeds={"sharma": 7, "kaveri": 11, "swift": 23}):
    made = {}
    for k in which:
        spec = SPECS[k]
        spec["bs_accounts"] = COMMON_BS | {spec["bank"], spec["cc"], spec["recv_acct"], spec["pay_acct"], spec["gst_in"], spec["gst_out"]} | \
            ({spec["tl"]} if spec.get("tl") else set()) | ({spec["inv_acct"]} if spec.get("inv_acct") else set()) | set(spec["fa0"])
        bk = simulate(spec, seeds[k])
        df = export(bk, spec)
        made[k] = df
        party = (spec["recv_acct"], spec["pay_acct"])
        if k == "sharma":
            write_tally(df, OUT / "Sharma_Traders_daybook_FY24-FY26.csv", spec["company"], party)
            write_by_account(df, OUT / "Sharma_Traders_ledger_by_account.xlsx", spec["company"], party)
        elif k == "kaveri":
            write_zoho(df, OUT / "Kaveri_Foods_journal_FY24-FY26.xlsx", spec["company"])
        else:
            write_messy(df, OUT / "Swift_Freight_journal_messy.xlsx", spec["company"])
        print(k, len(df), "lines", round(df.debit.sum() / 1e7, 1), "cr turnover of postings")
    return made


if __name__ == "__main__":
    build(tuple(sys.argv[1:]) or ("sharma", "kaveri", "swift"))
